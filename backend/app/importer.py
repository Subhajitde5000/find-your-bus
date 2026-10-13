"""Load the two Excel workbooks in /data into the PostgreSQL schema.

1. Kolkata_Bus_Stands_v2.xlsx                      -> operators-free `bus_stands`
2. Kolkata_Comprehensive_Bus_Routes_and_Fleet_Directory.xlsx
     - "Master Index & Directory"                   -> fleet allocation per route
     - one sheet per route (Route_*)                -> route header, stops, trips, buses

Run through `python manage.py import-data`. The import is all-or-nothing (one transaction).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime, time
from pathlib import Path

import openpyxl
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.models import (
    Bus,
    BusStand,
    Operator,
    Route,
    RouteStop,
    ServiceCategory,
    Stop,
    Trip,
    TripSchedule,
)
from app.stand_aliases import STAND_ALIASES

TABLES_TO_REPLACE = (
    "trip_schedules",
    "trips",
    "buses",
    "route_stops",
    "routes",
    "stops",
    "bus_stands",
    "service_categories",
    "operators",
)

_VIA_NOTE = re.compile(r"\s*\(via [^)]*\)\s*$", re.IGNORECASE)
_TIME_FMT = "%I:%M %p"


@dataclass
class ImportReport:
    stands: int = 0
    stops: int = 0
    stops_linked_to_stands: int = 0
    operators: int = 0
    service_categories: int = 0
    routes: int = 0
    route_stops: int = 0
    buses: int = 0
    trips: int = 0
    trip_schedules: int = 0
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> str:
        lines = [
            f"stands:              {self.stands}",
            f"operators:           {self.operators}",
            f"service categories:  {self.service_categories}",
            f"stops:               {self.stops} ({self.stops_linked_to_stands} linked to a stand)",
            f"routes:              {self.routes}",
            f"route stops:         {self.route_stops}",
            f"buses:               {self.buses}",
            f"trips:               {self.trips}",
            f"trip schedules:      {self.trip_schedules}",
        ]
        lines += [f"WARNING: {w}" for w in self.warnings]
        return "\n".join(lines)


def clean_text(value: object) -> str | None:
    if value is None:
        return None
    s = re.sub(r"\s+", " ", str(value)).strip()
    return s or None


def clean_stop_name(raw: str) -> str:
    """'Howrah Railway Station (via E.M. Bypass)' -> 'Howrah Railway Station'."""
    return _VIA_NOTE.sub("", raw).strip()


def normalise_registration(reg: str) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", reg).upper()


def parse_clock(value: str) -> time:
    return datetime.strptime(value.strip(), _TIME_FMT).time()


def parse_fleet_count(text_value: str | None) -> int | None:
    if not text_value:
        return None
    m = re.match(r"^(\d+)\b", text_value.strip())
    return int(m.group(1)) if m else None


# --------------------------------------------------------------------------- stands
def read_stands(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        if not isinstance(row[0], int):
            continue  # title, header and blank rows
        serial, name, local, category, division, address, role = row[:7]
        rows.append(
            {
                "serial_no": serial,
                "name": clean_text(name),
                "local_name": clean_text(local),
                "node_category": clean_text(category),
                "operator_division": clean_text(division),
                "address": clean_text(address),
                "functional_role": clean_text(role),
            }
        )
    return rows


# --------------------------------------------------------------------------- routes
def read_master_index(wb: openpyxl.Workbook) -> dict[str, dict]:
    ws = wb.worksheets[0]
    index: dict[str, dict] = {}
    for row in ws.iter_rows(values_only=True):
        if not isinstance(row[0], int) or not row[1]:
            continue
        route_no = clean_text(row[1])
        index[route_no] = {
            "origin": clean_text(row[3]),
            "destination": clean_text(row[4]),
            "fleet_text": clean_text(row[6]),
        }
    return index


def read_route_sheet(ws) -> dict:
    rows = list(ws.iter_rows(values_only=True))
    title = clean_text(rows[0][0]) or ""
    route_no = title.split(":", 1)[1].strip() if ":" in title else ws.title
    header = clean_text(rows[1][0]) or ""
    m = re.match(r"Operator:\s*(.*?)\s*\|\s*Category:\s*(.*)$", header)
    if not m:
        raise ValueError(f"{ws.title}: unexpected operator/category line: {header!r}")

    header_row = next(i for i, r in enumerate(rows) if r and r[0] == "Run / Trip ID")
    trips = []
    for r in rows[header_row + 1 :]:
        if not r or not r[0]:
            continue
        trips.append(
            {
                "run_code": clean_text(r[0]),
                "bus": clean_text(r[1]),
                "out_dep": parse_clock(r[2]),
                "out_arr": parse_clock(r[3]),
                "ret_dep": parse_clock(r[4]),
                "ret_arr": parse_clock(r[5]),
                "notes": clean_text(r[6]),
            }
        )

    return {
        "route_number": route_no,
        "sheet": ws.title,
        "operator": clean_text(m.group(1)),
        "category": clean_text(m.group(2)),
        "origin": clean_text(rows[3][1]),
        "destination": clean_text(rows[3][4]),
        "via": clean_text(rows[4][1]),
        "trips": trips,
    }


def split_via(via: str | None) -> list[str]:
    if not via:
        return []
    return [p for p in (clean_text(x) for x in via.split(",")) if p]


# --------------------------------------------------------------------------- main
class _StopCache:
    """Creates each distinct stop once and links it to a bus stand when possible."""

    def __init__(self, session: Session, stands_by_name: dict[str, BusStand],
                 stands_by_local: dict[str, BusStand], report: ImportReport):
        self.session = session
        self.by_key: dict[str, Stop] = {}
        self.stands_by_name = stands_by_name
        self.stands_by_local = stands_by_local
        self.report = report

    def get(self, raw_name: str) -> Stop:
        name = clean_stop_name(raw_name)
        key = name.lower()
        if key in self.by_key:
            return self.by_key[key]

        stand = self.stands_by_name.get(key) or self.stands_by_local.get(key)
        if stand is None and key in STAND_ALIASES:
            target = STAND_ALIASES[key]
            stand = self.stands_by_name.get(target.lower())
            if stand is None:
                raise ValueError(f"stand alias {key!r} points to unknown stand {target!r}")

        stop = Stop(name=name, bus_stand=stand)
        self.session.add(stop)
        self.by_key[key] = stop
        self.report.stops += 1
        if stand is not None:
            self.report.stops_linked_to_stands += 1
        return stop


def import_all(session: Session, stands_path: Path, routes_path: Path,
           replace: bool = False) -> ImportReport:
    report = ImportReport()

    existing = session.scalar(select(func.count()).select_from(BusStand))
    if existing and not replace:
        raise RuntimeError(
            "Database already contains data. Re-run with --replace to clear and reload it."
        )
    if replace:
        session.execute(
            text(f"TRUNCATE {', '.join(TABLES_TO_REPLACE)} RESTART IDENTITY CASCADE")
        )

    # 1. Stands -----------------------------------------------------------
    stand_rows = read_stands(stands_path)
    stands_by_name: dict[str, BusStand] = {}
    for row in stand_rows:
        stand = BusStand(**row)
        session.add(stand)
        stands_by_name[row["name"].lower()] = stand
    session.flush()
    report.stands = len(stand_rows)
    stands_by_local = {
        s.local_name.lower(): s for s in stands_by_name.values() if s.local_name
    }

    # 2. Routes workbook ---------------------------------------------------
    wb = openpyxl.load_workbook(routes_path, data_only=True)
    index = read_master_index(wb)
    sheets = [read_route_sheet(ws) for ws in wb.worksheets[1:]]

    operators: dict[str, Operator] = {}
    categories: dict[str, ServiceCategory] = {}
    stops = _StopCache(session, stands_by_name, stands_by_local, report)
    seen_routes: set[str] = set()

    for sheet in sheets:
        rno = sheet["route_number"]
        if rno in seen_routes:
            raise ValueError(f"duplicate route number {rno}")
        seen_routes.add(rno)

        entry = index.get(rno)
        if entry is None:
            report.warnings.append(f"route {rno} ({sheet['sheet']}) is not in the master index")
            fleet_text = None
        else:
            fleet_text = entry["fleet_text"]
            if (entry["origin"], entry["destination"]) != (sheet["origin"], sheet["destination"]):
                report.warnings.append(f"route {rno}: origin/destination differ from master index")

        op = operators.get(sheet["operator"])
        if op is None:
            op = operators[sheet["operator"]] = Operator(name=sheet["operator"])
            session.add(op)
        cat = categories.get(sheet["category"])
        if cat is None:
            cat = categories[sheet["category"]] = ServiceCategory(name=sheet["category"])
            session.add(cat)

        origin_stop = stops.get(sheet["origin"])
        destination_stop = stops.get(sheet["destination"])
        route = Route(
            route_number=rno,
            operator=op,
            service_category=cat,
            origin_stop=origin_stop,
            destination_stop=destination_stop,
            origin_label=sheet["origin"],
            destination_label=sheet["destination"],
            via_text=sheet["via"],
            fleet_allocation_text=fleet_text,
            documented_fleet_count=parse_fleet_count(fleet_text),
            source_sheet=sheet["sheet"],
        )
        session.add(route)
        report.routes += 1

        ordered = [origin_stop, *(stops.get(v) for v in split_via(sheet["via"])), destination_stop]
        for seq, stop in enumerate(ordered, start=1):
            session.add(RouteStop(route=route, sequence=seq, stop=stop))
            report.route_stops += 1
        session.flush()

        # Trips and buses -------------------------------------------------
        for t in sheet["trips"]:
            bus = None
            if t["bus"]:
                key = normalise_registration(t["bus"])
                bus = session.scalar(select(Bus).where(Bus.registration_key == key))
                if bus is None:
                    bus = Bus(registration_no=t["bus"], registration_key=key)
                    session.add(bus)
                    session.flush()
                    report.buses += 1

            trip = Trip(run_code=t["run_code"], route=route, bus=bus, notes=t["notes"])
            session.add(trip)
            session.add_all(
                [
                    TripSchedule(trip=trip, direction="outbound",
                                 departure_time=t["out_dep"], arrival_time=t["out_arr"]),
                    TripSchedule(trip=trip, direction="return",
                                 departure_time=t["ret_dep"], arrival_time=t["ret_arr"]),
                ]
            )
            report.trips += 1
            report.trip_schedules += 2

        session.flush()

    report.operators = len(operators)
    report.service_categories = len(categories)
    session.commit()
    return report
