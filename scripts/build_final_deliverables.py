#!/usr/bin/env python3
"""Build final_bus_rout_timntable.xlsx (all 651 verified routes) and
finel_busstop_location.csv (all 2,194 canonical bus stops) from the verified
website dataset, stop coordinate registry, and conflict audit logs.
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
WEB_JSON = ROOT / "fontend" / "src" / "data" / "kolkata-bus-data.json"
COORDS_JSON = ROOT / "data" / "stop_coordinates.json"
V8_CSV = ROOT / "stop_name_loactio-v8.csv"
STOP_AUDIT_CSV = ROOT / "change_stop_location.csv"
SHIFTED_CSV_LINES = set(range(255, 460)) | set(range(1463, 1609)) | {1641, 1283}

FINAL_XLSX = ROOT / "final_bus_rout_timntable.xlsx"
FINAL_CSV = ROOT / "finel_busstop_location.csv"


def fmt_12h(clock24: str | None) -> str | None:
    """Convert 'HH:MM' 24-hour string into 'HH:MM AM/PM' matching v28 style."""
    if not clock24:
        return None
    m = re.match(r"^(\d{1,2}):(\d{2})$", clock24.strip())
    if not m:
        return clock24
    hh, mm = int(m.group(1)) % 24, int(m.group(2))
    suf = "AM" if hh < 12 else "PM"
    h12 = hh % 12
    if h12 == 0:
        h12 = 12
    return f"{h12:02d}:{mm:02d} {suf}"


def make_sheet_title(route_number: str, used_titles: set[str]) -> str:
    """Create a valid, unique Excel worksheet title (<= 31 chars)."""
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", route_number.strip())
    cleaned = re.sub(r"\s+", " ", cleaned)
    base = f"Route_{cleaned}"
    if len(base) > 31:
        base = f"R_{cleaned}"[:31].rstrip()
    cand = base
    idx = 2
    while cand.lower() in used_titles:
        suffix = f"_{idx}"
        cand = f"{base[: 31 - len(suffix)]}{suffix}"
        idx += 1
    used_titles.add(cand.lower())
    return cand


def build_final_stop_csv(
    web_data: dict[str, Any],
    coords_data: dict[str, Any],
) -> None:
    stops_list: list[dict[str, Any]] = web_data["stops"]
    sc_stops: dict[str, dict[str, Any]] = coords_data.get("stops", {})

    # Load original v8 non-shifted rows by stop name
    v8_by_stop: dict[str, dict[str, str]] = {}
    with open(V8_CSV, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for line_no, row in enumerate(reader, start=2):
            sname = re.sub(r"\s+", " ", (row.get("stop") or "").strip())
            if sname and line_no not in SHIFTED_CSV_LINES and sname not in v8_by_stop:
                v8_by_stop[sname] = row

    # Load final audit record per stop
    audit_by_stop: dict[str, dict[str, str]] = {}
    if STOP_AUDIT_CSV.exists():
        with open(STOP_AUDIT_CSV, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                sname = (row.get("stop_name") or "").strip()
                if sname:
                    audit_by_stop[sname] = row

    out_rows: list[dict[str, str]] = []
    for stop_entry in stops_list:
        sname = stop_entry["name"]
        sc_info = sc_stops.get(sname)
        v8_row = v8_by_stop.get(sname)
        aud = audit_by_stop.get(sname)

        if sc_info and sc_info.get("lat") is not None and sc_info.get("lng") is not None:
            lat_str = f"{float(sc_info['lat']):.5f}".rstrip("0").rstrip(".")
            lng_str = f"{float(sc_info['lng']):.5f}".rstrip("0").rstrip(".")
            conf = sc_info.get("confidence") or "high"
            matched_name = sc_info.get("matchedName") or (v8_row.get("matched_name") if v8_row else "") or sname
            matched_on = sc_info.get("source") or (v8_row.get("matched_on") if v8_row else "") or "verified"
            alts = (v8_row.get("alternatives") if v8_row else "") or ""
            if aud and aud.get("new_lat") and aud.get("new_lng"):
                verdict = aud.get("conflict_type") or "verified"
            elif v8_row and v8_row.get("verdict"):
                verdict = v8_row["verdict"]
            else:
                verdict = "verified_corridor_stop"
            out_rows.append(
                {
                    "stop": sname,
                    "lat": lat_str,
                    "lng": lng_str,
                    "confidence": conf,
                    "matched_name": matched_name,
                    "matched_on": matched_on,
                    "alternatives": alts,
                    "verdict": verdict,
                }
            )
        else:
            verdict = (aud.get("conflict_type") if aud else None) or "unlocated_no_published_coord"
            out_rows.append(
                {
                    "stop": sname,
                    "lat": "",
                    "lng": "",
                    "confidence": "",
                    "matched_name": "",
                    "matched_on": "verified_unlocated",
                    "alternatives": "",
                    "verdict": verdict,
                }
            )

    fieldnames = [
        "stop",
        "lat",
        "lng",
        "confidence",
        "matched_name",
        "matched_on",
        "alternatives",
        "verdict",
    ]
    with open(FINAL_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(out_rows)

    located_cnt = sum(1 for r in out_rows if r["lat"] and r["lng"])
    print(
        f"Wrote {FINAL_CSV.name}: {len(out_rows)} total stops "
        f"({located_cnt} located, {len(out_rows) - located_cnt} unlocated)"
    )


def build_final_routes_xlsx(web_data: dict[str, Any]) -> None:
    routes: list[dict[str, Any]] = web_data["routes"]
    buses: list[dict[str, Any]] = web_data.get("buses", [])

    wb = openpyxl.Workbook()
    ws_idx = wb.active
    ws_idx.title = "Master Index & Directory"

    # Styles matching missing_timtable_rout-v28.xlsx
    font_title_idx = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_sub_idx = Font(name="Calibri", size=11, bold=False, color="FFFFFF")
    font_card_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_card_val = Font(name="Calibri", size=16, bold=True, color="2F5496")
    font_tbl_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=11, bold=False)
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_link = Font(name="Calibri", size=11, bold=False, color="0563C1", underline="single")
    font_back_link = Font(name="Calibri", size=11, bold=True, color="002060", underline="single")
    font_sheet_title = Font(name="Calibri", size=14, bold=True, color="1F4E78")

    fill_navy = PatternFill(fill_type="solid", start_color="2F5496", end_color="2F5496")
    fill_blue = PatternFill(fill_type="solid", start_color="4472C4", end_color="4472C4")
    fill_zebra = PatternFill(fill_type="solid", start_color="F2F2F2", end_color="F2F2F2")
    fill_reg = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")

    thin_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    # 1. Master Index Header & Summary Cards
    total_routes = len(routes)
    last_idx_row = 7 + total_routes

    ws_idx.merge_cells("A1:G1")
    ws_idx["A1"] = "Kolkata Master Public & Private Bus Route Directory"
    ws_idx["A1"].font = font_title_idx
    ws_idx["A1"].fill = fill_navy
    ws_idx["A1"].alignment = align_center

    ws_idx.merge_cells("A2:G2")
    ws_idx["A2"] = f"Comprehensive Verified Timetable & Route Repository | Total Documented Routes: {total_routes}"
    ws_idx["A2"].font = font_sub_idx
    ws_idx["A2"].fill = fill_navy
    ws_idx["A2"].alignment = align_center

    ws_idx.merge_cells("A4:C4")
    ws_idx["A4"] = "TOTAL ROUTE SHEETS"
    ws_idx["A4"].font = font_card_hdr
    ws_idx["A4"].fill = fill_blue
    ws_idx["A4"].alignment = align_center

    ws_idx.merge_cells("D4:E4")
    ws_idx["D4"] = "PRIVATE SYNDICATE ROUTES"
    ws_idx["D4"].font = font_card_hdr
    ws_idx["D4"].fill = fill_blue
    ws_idx["D4"].alignment = align_center

    ws_idx.merge_cells("F4:G4")
    ws_idx["F4"] = "GOVT / STATE ROUTES"
    ws_idx["F4"].font = font_card_hdr
    ws_idx["F4"].fill = fill_blue
    ws_idx["F4"].alignment = align_center

    ws_idx.merge_cells("A5:C5")
    ws_idx["A5"] = f"=COUNTA(B8:B{last_idx_row})"
    ws_idx["A5"].font = font_card_val
    ws_idx["A5"].alignment = align_center

    ws_idx.merge_cells("D5:E5")
    ws_idx["D5"] = (
        f'=COUNTIF(D8:D{last_idx_row}, "*Private*") + COUNTIF(D8:D{last_idx_row}, "*Minibus*") '
        f'+ COUNTIF(D8:D{last_idx_row}, "*Suburban*") + COUNTIF(D8:D{last_idx_row}, "*Regional*") '
        f'+ COUNTIF(D8:D{last_idx_row}, "*Joint*")'
    )
    ws_idx["D5"].font = font_card_val
    ws_idx["D5"].alignment = align_center

    ws_idx.merge_cells("F5:G5")
    ws_idx["F5"] = (
        f'=COUNTIF(D8:D{last_idx_row}, "*WBTC*") + COUNTIF(D8:D{last_idx_row}, "*SBSTC*") '
        f'+ COUNTIF(D8:D{last_idx_row}, "*NBSTC*") + COUNTIF(D8:D{last_idx_row}, "*CTC*") '
        f'+ COUNTIF(D8:D{last_idx_row}, "*CSTC*") + COUNTIF(D8:D{last_idx_row}, "*WBSTC*") '
        f'+ COUNTIF(D8:D{last_idx_row}, "*State*") + COUNTIF(D8:D{last_idx_row}, "*Express*") '
        f'+ COUNTIF(D8:D{last_idx_row}, "*Government*")'
    )
    ws_idx["F5"].font = font_card_val
    ws_idx["F5"].alignment = align_center

    idx_headers = [
        "SL No.",
        "Route Number",
        "Sheet Tab Link",
        "Operator / Service Category",
        "Origin Terminal (X)",
        "Destination Terminal (Y)",
        "Key Intermediate Alignment",
    ]
    for col_i, h_text in enumerate(idx_headers, start=1):
        c = ws_idx.cell(7, col_i, h_text)
        c.font = font_tbl_hdr
        c.fill = fill_navy
        c.alignment = align_center if col_i == 1 else align_left

    ws_idx.row_dimensions[1].height = 28.0
    ws_idx.row_dimensions[2].height = 20.0
    ws_idx.row_dimensions[4].height = 18.0
    ws_idx.row_dimensions[5].height = 24.0
    ws_idx.row_dimensions[7].height = 24.0

    idx_col_widths = {"A": 10.0, "B": 30.0, "C": 32.0, "D": 35.0, "E": 30.0, "F": 30.0, "G": 60.0}
    for col_letter, width in idx_col_widths.items():
        ws_idx.column_dimensions[col_letter].width = width

    used_titles: set[str] = {ws_idx.title.lower()}

    for sl_no, route in enumerate(routes, start=1):
        rnum = route["number"]
        sheet_title = make_sheet_title(rnum, used_titles)
        category = route.get("category") or route.get("operator") or "Private Stage Carriage"
        origin = route["origin"]
        destination = route["destination"]
        stops: list[str] = route["stops"]
        intermediate_stops = stops[1:-1] if len(stops) > 2 else []
        inter_str = ", ".join(intermediate_stops) if intermediate_stops else f"{origin}, {destination}"

        # Populate Master Index row
        idx_row = 7 + sl_no
        ws_idx.row_dimensions[idx_row].height = 20.0
        is_even = sl_no % 2 == 0
        row_fill = fill_zebra if is_even else None

        c_sl = ws_idx.cell(idx_row, 1, sl_no)
        c_rn = ws_idx.cell(idx_row, 2, rnum)
        c_lk = ws_idx.cell(idx_row, 3, f"Go to {rnum} Sheet →")
        c_lk.hyperlink = f"#'{sheet_title}'!A1"
        c_cat = ws_idx.cell(idx_row, 4, category)
        c_org = ws_idx.cell(idx_row, 5, origin)
        c_dst = ws_idx.cell(idx_row, 6, destination)
        c_int = ws_idx.cell(idx_row, 7, inter_str)

        for col_i, cell in enumerate((c_sl, c_rn, c_lk, c_cat, c_org, c_dst, c_int), start=1):
            cell.font = font_link if col_i == 3 else font_regular
            cell.alignment = align_center if col_i == 1 else align_left
            cell.border = border_cell
            if row_fill:
                cell.fill = row_fill

        # Create Route Worksheet
        ws = wb.create_sheet(title=sheet_title)
        ws.cell(1, 1, f"ROUTE TIMETABLE DIRECTORY: BUS ROUTE {rnum}").font = font_sheet_title

        ws.cell(3, 1, "Route Number:").font = font_bold
        ws.cell(3, 2, rnum).font = font_bold
        ws.cell(3, 4, "Category:").font = font_bold
        ws.cell(3, 5, category).font = font_regular

        ws.cell(4, 1, "Origin (Stand X):").font = font_bold
        ws.cell(4, 2, origin).font = font_regular
        ws.cell(4, 4, "Destination (Stand Y):").font = font_bold
        ws.cell(4, 5, destination).font = font_regular
        c_ret = ws.cell(4, 7, '=HYPERLINK("#\'Master Index & Directory\'!A1", "← Return to Master Index")')
        c_ret.font = font_back_link

        ws.cell(5, 1, "Intermediate Alignment:").font = font_bold
        ws.cell(5, 2, inter_str).font = font_regular

        out_runs = [r for r in route["runs"] if r["direction"] == "outbound"]
        ret_runs = [r for r in route["runs"] if r["direction"] == "return"]
        has_notes = any(bool(r.get("notes")) for r in route["runs"])

        def write_direction_section(
            start_row: int,
            sec_title: str,
            dir_stops: list[str],
            dir_runs: list[dict[str, Any]],
        ) -> int:
            c_banner = ws.cell(start_row, 1, sec_title)
            c_banner.font = font_card_hdr
            c_banner.fill = fill_blue

            hdr_row = start_row + 1
            headers = ["Trip ID", "Vehicle Reg No."]
            for s_idx, s_name in enumerate(dir_stops):
                if s_idx == 0:
                    headers.append(f"Departure ({s_name})")
                elif s_idx == len(dir_stops) - 1:
                    headers.append(f"Arrival ({s_name})")
                else:
                    headers.append(f"Passage {s_idx} ({s_name})")
            if has_notes:
                headers.append("Operational Notes")

            for col_i, h_val in enumerate(headers, start=1):
                hc = ws.cell(hdr_row, col_i, h_val)
                hc.font = font_tbl_hdr
                hc.fill = fill_navy
                hc.alignment = align_left

            cur_r = hdr_row + 1
            for run_idx, run_obj in enumerate(dir_runs, start=1):
                bus_idx = run_obj.get("bus")
                if bus_idx is not None and 0 <= bus_idx < len(buses):
                    b_entry = buses[bus_idx]
                    reg_no = (b_entry.get("reg") if isinstance(b_entry, dict) else str(b_entry)) or "Unindexed (Private Roster)"
                else:
                    reg_no = "Unindexed (Private Roster)"

                z_fill = fill_zebra if (run_idx % 2 == 0) else None
                c_tid = ws.cell(cur_r, 1, run_obj["run"])
                c_tid.font = font_regular
                if z_fill:
                    c_tid.fill = z_fill

                c_reg = ws.cell(cur_r, 2, reg_no)
                c_reg.font = font_regular
                c_reg.fill = fill_reg

                times_list = run_obj.get("times") or []
                for s_idx in range(len(dir_stops)):
                    raw_t = times_list[s_idx] if s_idx < len(times_list) else None
                    tc = ws.cell(cur_r, 3 + s_idx, fmt_12h(raw_t))
                    tc.font = font_regular
                    if z_fill:
                        tc.fill = z_fill

                if has_notes:
                    nc = ws.cell(cur_r, 3 + len(dir_stops), run_obj.get("notes") or "")
                    nc.font = font_regular
                    if z_fill:
                        nc.fill = z_fill

                cur_r += 1

            return cur_r

        next_r = write_direction_section(
            7,
            f"OUTBOUND TRIP TIMETABLE ({origin} → {destination})",
            stops,
            out_runs,
        )

        ret_start = next_r + 2
        rev_stops = list(reversed(stops))
        end_r = write_direction_section(
            ret_start,
            f"RETURN TRIP TIMETABLE ({destination} → {origin})",
            rev_stops,
            ret_runs,
        )

        total_cols = 2 + len(stops) + (1 if has_notes else 0)
        for col_i in range(1, total_cols + 1):
            col_letter = get_column_letter(col_i)
            ws.column_dimensions[col_letter].width = 22.0 if col_i in (1, 3) else 26.0
        _ = end_r

    wb.save(FINAL_XLSX)
    print(f"Wrote {FINAL_XLSX.name}: {len(wb.sheetnames)} total sheets (1 Master Index + {len(routes)} route sheets)")


def main() -> int:
    web_data = json.loads(WEB_JSON.read_text(encoding="utf-8"))
    coords_data = json.loads(COORDS_JSON.read_text(encoding="utf-8"))

    build_final_stop_csv(web_data, coords_data)
    build_final_routes_xlsx(web_data)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
